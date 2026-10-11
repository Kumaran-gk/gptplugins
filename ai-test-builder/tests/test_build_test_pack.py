from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document

from src.agent.workflow import AgentWorkflow, WorkflowRequest
from src.utils.json_utils import load_json_file, write_json_file


def test_build_test_pack_runs_analysis_and_multiple_generation_versions(temp_settings):
    subject = "chemistry"
    subject_dirs = temp_settings.ensure_subject_dirs(subject)
    input_path = Path(subject_dirs["input_path"]) / "unit-2-test.docx"

    document = Document()
    document.add_paragraph("SCH3U Unit 2 Chemical Reactions Test")
    document.add_paragraph("Section A - Multiple Choice")
    document.add_paragraph("1. Which reaction is a synthesis reaction? (1 mark)")
    document.add_paragraph("2. Predict the products for magnesium reacting with oxygen. (2 marks)")
    document.add_paragraph("Section B - Written Response")
    document.add_paragraph("3. Balance and classify: Al + O2 -> (3 marks)")
    document.add_paragraph("4. Explain one application of combustion reactions. (3 marks)")
    document.save(str(input_path))

    workflow = AgentWorkflow(settings=temp_settings)
    request = WorkflowRequest(
        skill="build-test-pack",
        subject=subject,
        input_files=[str(input_path)],
        options={
            "versions": "verA,verB",
            "include_teacher": "true",
            "save_debug": "false",
            "export_formats": "md,docx,pdf",
            "manifest_label": "demo-pack",
        },
    )

    result = workflow.run(request)

    assert result["status"] == "ok"
    assert result["analysis_result"] is not None
    assert len(result["blueprint_paths"]) == 1
    assert result["versions_requested"] == ["verA", "verB"]
    assert len(result["version_results"]) == 2
    assert len(result["generated_student_paths"]) == 2
    assert len(result["generated_teacher_paths"]) == 2
    assert len(result["generated_student_docx_paths"]) == 2
    assert len(result["generated_teacher_docx_paths"]) == 2
    assert len(result["generated_student_pdf_paths"]) == 2
    assert len(result["generated_teacher_pdf_paths"]) == 2
    assert result["cumulative_state"] == "not_requested"

    manifest_path = Path(result["manifest_path"])
    manifest = load_json_file(manifest_path)
    artifact_root = Path(subject_dirs["generated_path"]) / "artifacts"

    assert manifest_path.exists()
    assert artifact_root in manifest_path.parents
    assert manifest["versions_requested"] == ["verA", "verB"]
    assert len(manifest["generated_student_paths"]) == 2
    assert len(manifest["generated_student_docx_paths"]) == 2
    assert len(manifest["generated_student_pdf_paths"]) == 2
    assert manifest["cumulative_result"] is None

    for path in result["generated_student_docx_paths"] + result["generated_teacher_docx_paths"]:
        file_path = Path(path)
        assert file_path.exists()
        assert artifact_root in file_path.parents
    for path in result["generated_student_pdf_paths"] + result["generated_teacher_pdf_paths"]:
        file_path = Path(path)
        assert file_path.exists()
        assert file_path.parent == Path(subject_dirs["generated_path"])


def test_build_test_pack_can_build_weighted_cumulative_from_blueprints(
    temp_settings,
    valid_blueprint,
):
    subject = "chemistry"
    subject_dirs = temp_settings.ensure_subject_dirs(subject)

    first = deepcopy(valid_blueprint)
    first["blueprint_id"] = "chemistry-unit-2"
    first["title"] = "Unit 2 Chemical Reactions Test"
    first["unit"]["unit_name"] = "Unit 2 Chemical Reactions"
    first["topic_distribution"] = [
        {
            "topic": "chemical_reactions",
            "weight_pct": 70.0,
            "subtopics": ["balance", "combustion"],
        },
        {
            "topic": "reaction_rates",
            "weight_pct": 30.0,
            "subtopics": ["collision_theory"],
        },
    ]

    second = deepcopy(valid_blueprint)
    second["blueprint_id"] = "chemistry-unit-3"
    second["title"] = "Unit 3 Stoichiometry Test"
    second["unit"]["unit_name"] = "Unit 3 Stoichiometry"
    second["topic_distribution"] = [
        {
            "topic": "stoichiometry",
            "weight_pct": 60.0,
            "subtopics": ["mole", "yield"],
        },
        {
            "topic": "gas_laws",
            "weight_pct": 40.0,
            "subtopics": ["ideal_gas"],
        },
    ]

    first_path = Path(subject_dirs["blueprint_path"]) / "chemistry_unit-2_blueprint.json"
    second_path = Path(subject_dirs["blueprint_path"]) / "chemistry_unit-3_blueprint.json"
    write_json_file(first_path, first)
    write_json_file(second_path, second)

    workflow = AgentWorkflow(settings=temp_settings)
    request = WorkflowRequest(
        skill="build-test-pack",
        subject=subject,
        blueprint_files=[str(first_path), str(second_path)],
        options={
            "versions": "verPack",
            "include_teacher": "true",
            "save_debug": "false",
            "export_formats": "md,docx,pdf",
            "build_cumulative": "true",
            "cumulative_version": "verFinal",
            "unit_weights": "Unit 2 Chemical Reactions:25,Unit 3 Stoichiometry:75",
            "manifest_label": "weighted-pack",
        },
    )

    result = workflow.run(request)

    assert result["status"] == "ok"
    assert result["cumulative_state"] == "generated"
    assert result["cumulative_result"] is not None
    assert result["cumulative_result"]["unit_weights"] == {
        "Unit 2 Chemical Reactions": 25.0,
        "Unit 3 Stoichiometry": 75.0,
    }
    assert result["cumulative_result"]["cumulative_plan"]["merge_strategy"] == (
        "weighted_unit_topic_round_robin"
    )

    top_topics = result["cumulative_result"]["cumulative_plan"]["top_topics"]
    assert top_topics.index("gas_laws") < top_topics.index("reaction_rates")

    student_path = Path(result["cumulative_result"]["cumulative_student_test_path"])
    student_text = student_path.read_text(encoding="utf-8")
    student_docx_path = Path(result["cumulative_result"]["cumulative_student_docx_path"])
    teacher_docx_path = Path(result["cumulative_result"]["cumulative_teacher_docx_path"])
    student_pdf_path = Path(result["cumulative_result"]["cumulative_student_pdf_path"])
    teacher_pdf_path = Path(result["cumulative_result"]["cumulative_teacher_pdf_path"])

    assert "**Unit Weighting:** Unit 2 Chemical Reactions: 25%, Unit 3 Stoichiometry: 75%" in student_text
    assert student_docx_path.exists()
    assert teacher_docx_path.exists()
    assert student_pdf_path.exists()
    assert teacher_pdf_path.exists()
    artifact_root = Path(subject_dirs["generated_path"]) / "artifacts"
    assert artifact_root in student_path.parents
    assert artifact_root in student_docx_path.parents
    assert artifact_root in teacher_docx_path.parents
    assert student_pdf_path.parent == Path(subject_dirs["generated_path"])
    assert teacher_pdf_path.parent == Path(subject_dirs["generated_path"])


def test_build_test_pack_can_ground_generation_with_unit_content(temp_settings):
    subject = "chemistry"
    subject_dirs = temp_settings.ensure_subject_dirs(subject)
    input_path = Path(subject_dirs["input_path"]) / "unit-2-test.docx"
    unit_content_path = Path(subject_dirs["input_path"]) / "unit-2-content.docx"

    test_document = Document()
    test_document.add_paragraph("SCH3U Unit 2 Chemical Reactions Test")
    test_document.add_paragraph("Section A - Multiple Choice")
    test_document.add_paragraph("1. Which reaction is a synthesis reaction? (1 mark)")
    test_document.add_paragraph("2. Predict the products for magnesium reacting with oxygen. (2 marks)")
    test_document.add_paragraph("Section B - Written Response")
    test_document.add_paragraph("3. Balance and classify: Al + O2 -> (3 marks)")
    test_document.add_paragraph("4. Explain one application of combustion reactions. (3 marks)")
    test_document.save(str(input_path))

    unit_document = Document()
    unit_document.add_paragraph(
        "Argentite ore contains silver sulfide, Ag2S, and can be processed through displacement ideas to isolate silver."
    )
    unit_document.add_paragraph(
        "Calcium carbonate neutralizes hydrochloric acid while forming calcium chloride, water, and carbon dioxide."
    )
    unit_document.save(str(unit_content_path))

    workflow = AgentWorkflow(settings=temp_settings)
    request = WorkflowRequest(
        skill="build-test-pack",
        subject=subject,
        input_files=[str(input_path)],
        unit_content_files=[str(unit_content_path)],
        options={
            "versions": "verGrounded",
            "include_teacher": "false",
            "save_debug": "false",
            "export_formats": "md",
            "manifest_label": "grounded-pack",
        },
    )

    result = workflow.run(request)

    assert result["status"] == "ok"
    assert result["unit_content_bundle_path"] is not None
    assert result["unit_content_text_path"] is not None
    assert result["unit_content_summary"]["snippet_count"] > 0

    student_path = Path(result["generated_student_paths"][0])
    student_text = student_path.read_text(encoding="utf-8")

    assert Path(result["unit_content_bundle_path"]).exists()
    assert Path(result["unit_content_text_path"]).exists()
    assert "Argentite ore contains silver sulfide" in student_text
    artifact_root = Path(subject_dirs["generated_path"]) / "artifacts"
    assert artifact_root in Path(result["unit_content_bundle_path"]).parents
    assert artifact_root in Path(result["unit_content_text_path"]).parents
    assert artifact_root in Path(result["generated_student_paths"][0]).parents
