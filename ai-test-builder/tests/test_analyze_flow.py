from __future__ import annotations

import json
from pathlib import Path

from docx import Document

from src.agent.workflow import AgentWorkflow, WorkflowRequest


def test_analyze_flow_creates_blueprint_and_memory(temp_settings):
    subject = "chemistry"
    subject_dirs = temp_settings.ensure_subject_dirs(subject)
    input_path = Path(subject_dirs["input_path"]) / "unit-2-test.docx"

    document = Document()
    document.add_paragraph("SCH3U Unit 2 Chemical Reactions Test")
    document.add_paragraph("Section A - Multiple Choice")
    document.add_paragraph("1. Which reaction is a synthesis reaction? (1 mark)")
    document.add_paragraph("2. Predict the products for magnesium reacting with oxygen. (2 marks)")
    document.add_paragraph("Section B - Written Response")
    document.add_paragraph("3. Balance and classify: Al + O2 -> (3 marks)")
    document.add_paragraph("4. Explain one application of combustion reactions. (3 marks)")
    document.save(str(input_path))

    workflow = AgentWorkflow(settings=temp_settings)
    request = WorkflowRequest(
        skill="analyze-test",
        subject=subject,
        input_files=[str(input_path)],
        options={"save_debug": "true"},
    )

    result = workflow.run(request)

    assert result["status"] in {"ok", "partial"}
    assert len(result["extracted_files"]) == 1
    assert len(result["blueprint_outputs"]) == 1

    blueprint_output = result["blueprint_outputs"][0]
    assert blueprint_output["blueprint_valid"] is True
    assert blueprint_output["blueprint_path"] is not None
    assert Path(blueprint_output["blueprint_path"]).exists()

    memory_path = blueprint_output["memory_record_path"]
    assert memory_path is not None
    assert Path(memory_path).exists()

    blueprint = json.loads(Path(blueprint_output["blueprint_path"]).read_text(encoding="utf-8"))
    sections = blueprint["structure"]["sections"]

    assert len(sections) >= 2
    assert sections[0].get("num_questions") == 2
    assert sections[1].get("num_questions") == 2
