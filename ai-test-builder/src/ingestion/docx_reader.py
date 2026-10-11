from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List

from docx import Document

from src.utils.logging_utils import get_logger


LOGGER = get_logger("ingestion.docx_reader")


def _read_paragraphs(document: Document) -> List[Dict[str, Any]]:
    segments: List[Dict[str, Any]] = []

    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if text:
            segments.append(
                {
                    "segment_type": "paragraph",
                    "index": index,
                    "text": text,
                    "char_count": len(text),
                }
            )

    return segments


def _read_tables(document: Document, start_index: int) -> List[Dict[str, Any]]:
    segments: List[Dict[str, Any]] = []
    running_index = start_index

    for table in document.tables:
        row_lines: List[str] = []

        for row in table.rows:
            cell_values = [cell.text.strip() for cell in row.cells]
            cleaned = [value for value in cell_values if value]
            if cleaned:
                row_lines.append(" | ".join(cleaned))

        table_text = "\n".join(row_lines).strip()
        if table_text:
            segments.append(
                {
                    "segment_type": "table",
                    "index": running_index,
                    "text": table_text,
                    "char_count": len(table_text),
                }
            )
            running_index += 1

    return segments


def read_docx(file_path: str | Path) -> Dict[str, Any]:
    """
    Read a DOCX and return a structured extraction payload.
    """
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"DOCX file not found: {path}")

    if path.suffix.lower() != ".docx":
        raise ValueError(f"Expected a .docx file, got: {path.name}")

    LOGGER.info("Reading DOCX path=%s", path)
    document = Document(str(path))

    paragraph_segments = _read_paragraphs(document)
    table_segments = _read_tables(document, start_index=len(paragraph_segments) + 1)

    segments = paragraph_segments + table_segments
    full_text = "\n\n".join(segment["text"] for segment in segments).strip()

    result = {
        "file_name": path.name,
        "file_path": str(path.resolve()),
        "file_type": "docx",
        "source_kind": "docx",
        "page_count": None,
        "text": full_text,
        "segments": segments,
        "metadata": {
            "paragraph_count": len(document.paragraphs),
            "table_count": len(document.tables),
            "extraction_method": "python-docx",
        },
    }
    LOGGER.info(
        "Completed DOCX read path=%s segments=%s paragraphs=%s tables=%s",
        path,
        len(segments),
        len(document.paragraphs),
        len(document.tables),
    )
    return result
